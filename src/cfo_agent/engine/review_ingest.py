"""Read a reviewed workbook back in. Maps each Ledger row to its ledger line by
the hidden ID column, and reports the reviewer's final category + billable flag.
The CLI turns changes into approved decisions and new rules."""
from __future__ import annotations

from pathlib import Path

import openpyxl

PLACEHOLDER = "(reviewer to choose)"


def read_reviewed(path: Path) -> list:
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Ledger" not in wb.sheetnames:
        raise ValueError("No 'Ledger' sheet in the reviewed workbook.")
    ws = wb["Ledger"]
    header = [c.value for c in ws[1]]
    idx = {name: i for i, name in enumerate(header)}
    for needed in ("Proposed COA", "ID"):
        if needed not in idx:
            raise ValueError(f"Reviewed workbook missing '{needed}' column.")
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        ext = row[idx["ID"]]
        if not ext:
            continue
        coa = row[idx["Proposed COA"]]
        coa = (str(coa).strip() if coa is not None else "")
        if coa == PLACEHOLDER:
            coa = ""
        billable = None
        if "Billable" in idx and row[idx["Billable"]] is not None:
            billable = str(row[idx["Billable"]]).strip().upper() == "YES"
        out.append({"external_id": str(ext), "final_coa": coa, "billable": billable})
    return out
